import os
import sys
import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import reportlab
from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table as RLTable,
    TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

def clean_xml_text(text: str) -> str:
    """Sanitizes control characters to ensure XML compatibility for python-docx and ReportLab."""
    if not isinstance(text, str):
        return str(text)
    # Remove control characters except tab, newline, carriage return
    return "".join(ch for ch in text if ord(ch) >= 32 or ch in "\n\r\t")

# ----------------------------------------------------
# NUMBERED CANVAS FOR PDF PAGE NUMBERS (Page X of Y)
# ----------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_number(self, page_count):
        if self._pageNumber == 1:
            return  # Suppress page number on title page
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#4B5563"))
        
        # Header
        self.setLineWidth(0.5)
        self.setStrokeColor(colors.HexColor("#D1D5DB"))
        self.line(54, 750, 612 - 54, 750)
        self.drawString(54, 755, "Parcl Co. Ltd / Unified Mentor — Real Estate Buyer Segmentation Research Paper")
        
        # Footer
        self.line(54, 50, 612 - 54, 50)
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(612 - 54, 35, page_text)
        self.drawString(54, 35, "Confidential — Machine Learning Project Report")
        self.restoreState()


# ----------------------------------------------------
# 1. GENERATE DOCX FILE
# ----------------------------------------------------
def build_docx_paper(output_path="reports/real_estate_buyer_segmentation_research_paper.docx"):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    def add_custom_heading(text, level):
        text = clean_xml_text(text)
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        if level == 1:
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
        elif level == 2:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8) # Blue
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif level == 3:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51) # Dark Gray
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        return p

    def add_p(text, bold_prefix="", italic=False):
        text = clean_xml_text(text)
        bold_prefix = clean_xml_text(bold_prefix)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
        run = p.add_run(text)
        run.italic = italic
        return p

    def add_img(img_path, caption="", width=Inches(5.5)):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.add_run().add_picture(img_path, width=width)
            if caption:
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_after = Pt(10)
                crun = cp.add_run(clean_xml_text(caption))
                crun.italic = True
                crun.font.size = Pt(9.5)
                crun.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # TITLE PAGE (DOCX)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(72)
    p_title.paragraph_format.space_after = Pt(18)
    r_t = p_title.add_run("Machine Learning Based Buyer Segmentation and Investment Profiling for Real Estate Market Intelligence")
    r_t.bold = True
    r_t.font.size = Pt(24)
    r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(144)
    r_s = p_sub.add_run("A Data-Driven Unsupervised Machine Learning Framework & Interactive Market Intelligence Dashboard")
    r_s.italic = True
    r_s.font.size = Pt(13)
    r_s.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Domain:", "Financial Analytics & Real Estate Market Intelligence"),
        ("Internship Organization:", "Unified Mentor"),
        ("Collaborating Organization:", "Parcl Co. Limited"),
        ("Prepared By:", "Machine Learning Engineering Intern"),
        ("Submission Date:", "October 2026")
    ]
    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        r0 = row.cells[0].paragraphs[0].add_run(label)
        r0.bold = True
        r0.font.size = Pt(11)
        r1 = row.cells[1].paragraphs[0].add_run(val)
        r1.font.size = Pt(11)

    doc.add_page_break()

    # Generate Markdown First
    from generate_research_paper import generate_markdown_paper
    generate_markdown_paper()
    
    with open("reports/real_estate_buyer_segmentation_research_paper.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_title_page = True
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith("---") and in_title_page:
            in_title_page = False
            continue
        if in_title_page:
            continue

        if line_s.startswith("## "):
            title_text = line_s[3:].strip()
            add_custom_heading(title_text, level=1)
        elif line_s.startswith("### "):
            title_text = line_s[4:].strip()
            add_custom_heading(title_text, level=2)
        elif line_s.startswith("#### "):
            title_text = line_s[5:].strip()
            add_custom_heading(title_text, level=3)
        elif line_s.startswith("- "):
            add_p(line_s[2:].strip(), bold_prefix="- ")
        elif line_s.startswith("!["):
            try:
                caption = line_s[line_s.find("[")+1:line_s.find("]")]
                path_str = line_s[line_s.find("(")+1:line_s.find(")")]
                add_img(path_str, caption=caption)
            except Exception:
                pass
        elif not line_s.startswith("|"):
            add_p(line_s)

    # Insert Core Embedded Figures
    add_custom_heading("Embedded Figure Artifacts", level=2)
    add_img("reports/figures/elbow_method.png", "Figure 1: Elbow Method Plot (Inertia vs Number of Clusters K)")
    add_img("reports/figures/silhouette_scores.png", "Figure 2: Silhouette Score Plot across Tested Clusters (K=2 to 10)")
    add_img("reports/figures/cluster_distribution.png", "Figure 3: Buyer Distribution Share Across Clusters (K=4)")
    add_img("reports/figures/cluster_characteristics.png", "Figure 4: Cluster Characteristics Comparison Dashboard")
    add_img("reports/figures/sale_price_distribution.png", "Figure 5: Property Sale Price Distribution")
    add_img("reports/figures/buyer_age_distribution.png", "Figure 6: Buyer Age Distribution")

    doc.save(output_path)
    print(f"[SUCCESS] DOCX research paper saved to: {output_path}")


# ----------------------------------------------------
# 2. GENERATE PDF FILE (REPORTLAB)
# ----------------------------------------------------
def build_pdf_paper(output_path="reports/real_estate_buyer_segmentation_research_paper.pdf"):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    
    style_title = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1,
        spaceAfter=15
    )

    style_sub = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#4B5563'),
        alignment=1,
        spaceAfter=120
    )

    style_h1 = ParagraphStyle(
        'DocH1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )

    style_h2 = ParagraphStyle(
        'DocH2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1D4ED8'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    style_body = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1F2937'),
        spaceAfter=6
    )

    style_bullet = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        leftIndent=15,
        textColor=colors.HexColor('#1F2937'),
        spaceAfter=4
    )

    style_caption = ParagraphStyle(
        'DocCaption',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#4B5563'),
        spaceBefore=4,
        spaceAfter=10
    )

    story = []

    # TITLE PAGE
    story.append(Spacer(1, 40))
    story.append(Paragraph(clean_xml_text("Machine Learning Based Buyer Segmentation and Investment Profiling for Real Estate Market Intelligence"), style_title))
    story.append(Paragraph(clean_xml_text("A Data-Driven Unsupervised Machine Learning Framework & Interactive Market Intelligence Dashboard"), style_sub))

    meta_table_data = [
        [Paragraph("<b>Domain:</b>", style_body), Paragraph("Financial Analytics & Real Estate Market Intelligence", style_body)],
        [Paragraph("<b>Internship Organization:</b>", style_body), Paragraph("Unified Mentor", style_body)],
        [Paragraph("<b>Collaborating Organization:</b>", style_body), Paragraph("Parcl Co. Limited", style_body)],
        [Paragraph("<b>Prepared By:</b>", style_body), Paragraph("Machine Learning Engineering Intern", style_body)],
        [Paragraph("<b>Submission Date:</b>", style_body), Paragraph("October 2026", style_body)]
    ]
    t_meta = RLTable(meta_table_data, colWidths=[160, 320])
    t_meta.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_meta)
    story.append(PageBreak())

    # READ MARKDOWN CONTENT FOR PDF
    with open("reports/real_estate_buyer_segmentation_research_paper.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_title = True
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith("---") and in_title:
            in_title = False
            continue
        if in_title:
            continue

        # Format bold Markdown **text** -> ReportLab <b>text</b> cleanly
        formatted_text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', clean_xml_text(line_s))
        formatted_text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', formatted_text)
        
        if line_s.startswith("## "):
            story.append(Paragraph(clean_xml_text(line_s[3:].strip()), style_h1))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E5E7EB'), spaceBefore=2, spaceAfter=8))
        elif line_s.startswith("### "):
            story.append(Paragraph(clean_xml_text(line_s[4:].strip()), style_h2))
        elif line_s.startswith("- "):
            story.append(Paragraph(f"- {clean_xml_text(line_s[2:].strip())}", style_bullet))
        elif not line_s.startswith("|") and not line_s.startswith("!["):
            story.append(Paragraph(formatted_text, style_body))

    # ADD EMBEDDED FIGURES IN PDF
    story.append(PageBreak())
    story.append(Paragraph("Embedded Figure Artifacts & Visualizations", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E5E7EB'), spaceBefore=2, spaceAfter=8))

    fig_files = [
        ("reports/figures/elbow_method.png", "Figure 1: Elbow Method Plot (Inertia vs Number of Clusters K)"),
        ("reports/figures/silhouette_scores.png", "Figure 2: Silhouette Score Plot across Tested Clusters (K=2 to 10)"),
        ("reports/figures/cluster_distribution.png", "Figure 3: Buyer Distribution Share Across Clusters (K=4)"),
        ("reports/figures/cluster_characteristics.png", "Figure 4: Cluster Characteristics Comparison Dashboard"),
        ("reports/figures/sale_price_distribution.png", "Figure 5: Property Sale Price Distribution"),
        ("reports/figures/buyer_age_distribution.png", "Figure 6: Buyer Age Distribution")
    ]

    for fpath, cap in fig_files:
        if os.path.exists(fpath):
            img = RLImage(fpath, width=450, height=270)
            story.append(Spacer(1, 10))
            story.append(img)
            story.append(Paragraph(clean_xml_text(cap), style_caption))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[SUCCESS] PDF research paper saved to: {output_path}")

if __name__ == "__main__":
    build_docx_paper()
    build_pdf_paper()
